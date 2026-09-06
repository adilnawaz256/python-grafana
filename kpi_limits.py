import openpyxl
import re
import os
import json
import logging

logger = logging.getLogger("KPILimits")

EXCEL_PATH = "/Users/adilnawaz/Downloads/upper-lower limit for kpis.xlsx"

# DB Column Mapping for 59 NOC KPIs
DB_COLUMN_MAP = {
    'Latency': 'latency',
    'Packet Loss': 'packet_loss',
    'Links Down': 'links_down',
    'Transport Jitter': 'transport_jitter',
    'Transport Packet Loss': 'transport_packet_loss',
    'Backhaul Availability': 'backhaul_availability',
    'SMSC Traffic': 'smsc_traffic',
    'SMSC Resource Usage': 'smsc_resource_usage',
    'SMS Submission Success Rate': 'sms_submission_success_rate',
    'SMSC Availability': 'smsc_availability',
    'VoLTE CSSR': 'volte_cssr',
    'VoLTE Drop Rate': 'volte_drop_rate',
    'RRC Success Rate': 'rrc_success_rate',
    'E-RAB Success Rate': 'e_rab_success_rate',
    'Availability': 'availability',
    'S1 Success Rate': 's1_success_rate',
    'RAN Active Alarms': 'ran_active_alarms',
    'Site Down': 'site_down',
    'Intra-LTE HO Success Rate': 'intra_lte_ho_success_rate',
    'E-RAB Drop Rate': 'e_rab_drop_rate',
    'Critical Alarms Count': 'critical_alarms_count',
    'MTTR': 'mttr',
    'MCPTT Setup Time': 'mcptt_setup_time',
    'Emergency Call Success Rate': 'emergency_call_success_rate',
    'Online Users': 'online_users',
    'MC Push-To-Talk Attempts': 'mc_push_to_talk_attempts',
    'MC Push-To-Talk Success': 'mc_push_to_talk_success',
    'MC Video Attempts': 'mc_video_attempts',
    'MC Video Success': 'mc_video_success',
    'MC Data Attempts': 'mc_data_attempts',
    'MC Data Success': 'mc_data_success',
    'Total Number of LTE Connected Subs': 'total_number_of_lte_connected_subs',
    'Total Number of WiFi Connected Subs': 'total_number_of_wifi_connected_subs',
    'Total Traffic Volume Handled by MCX (UL+DL)': 'total_traffic_volume_handled_by_mcx_ul_dl',
    'MC VoIP Calls Audio Success': 'mc_voip_calls_audio_success',
    'IoT Attach Success Rate': 'iot_attach_success_rate',
    'IoT UL Message Success Rate': 'iot_ul_message_success_rate',
    'IoT DL Message Success Rate': 'iot_dl_message_success_rate',
    'IoT RACH Success Rate': 'iot_rach_success_rate',
    'IoT Device Availability': 'iot_device_availability',
    'SIP Registration SR': 'sip_registration_sr',
    'MO Session Connection Rate': 'mo_session_connection_rate',
    'MT Session Connection Rate': 'mt_session_connection_rate',
    'Call Setup Time (ms)': 'call_setup_time_ms',
    'Call Drop Rate of IMS': 'call_drop_rate_of_ims_mo',
    'Total Traffic (MO+MT) (Erl)': 'total_traffic_mo_mt_erl',
    'Registered Users': 'max_registered_users',
    'End-to-End Availability': 'end_to_end_availability',
    'SLA Compliance': 'sla_compliance',
    'DDoS Attacks Detected': 'ddos_attacks_detected',
    'Attach Success Rate': 'attach_success_rate',
    'Utilization': 'mme_cpu_utilization',
    'Service Request Success Rate': 'service_request_success_rate',
    'Paging Success Rate': 'paging_success_rate',
    'EPS Bearer Setup Success Rate': 'eps_bearer_setup_success_rate',
    'EPS Bearer Drop Rate': 'eps_bearer_drop_rate',
    'Provisioning Success Rate': 'provisioning_success_rate',
    'Network Trouble Tickets Open': 'network_trouble_tickets_open',
    'Customer Trouble Tickets Open': 'customer_trouble_tickets_open'
}

def parse_limit_rules(excel_path=EXCEL_PATH):
    if not os.path.exists(excel_path):
        logger.warning(f"Excel file not found at {excel_path}. Using fallback default rules.")
        return {}

    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws = wb.active

    kpi_rules = {}

    for r in range(2, ws.max_row + 1):
        name = str(ws.cell(r, 3).value or '').strip()
        upper = str(ws.cell(r, 5).value or '').strip()
        lower = str(ws.cell(r, 6).value or '').strip()

        col_name = DB_COLUMN_MAP.get(name)
        if not col_name:
            continue

        # Parse Upper (Healthy) and Lower (Warning) numeric ranges
        healthy_range = None
        warning_range = None

        # Greater than / equal percentage (e.g. >=99%, >=99.9%, >=99.5%)
        if '≥' in upper or '>=' in upper:
            val = float(re.search(r'[\d\.]+', upper).group())
            healthy_range = (val, 100.0)
            if '–' in lower or '-' in lower:
                nums = [float(x) for x in re.findall(r'[\d\.]+', lower)]
                if len(nums) == 2:
                    warning_range = (nums[0], nums[1])
            elif '≥' in lower or '>=' in lower:
                w_val = float(re.search(r'[\d\.]+', lower).group())
                warning_range = (w_val, val)

        # Less than limits (e.g. <30ms, <0.8%, <5ms, <2000ms, <75%)
        elif '<' in upper:
            val = float(re.search(r'[\d\.]+', upper).group())
            healthy_range = (max(0.0, val * 0.2), val)
            if '–' in lower or '-' in lower:
                nums = [float(x) for x in re.findall(r'[\d\.]+', lower)]
                if len(nums) == 2:
                    warning_range = (nums[0], nums[1])

        # Exact zero / counts (e.g. Links Down=0, Alarms=0, Site Down=0)
        elif upper == '0' or upper == '0 Critical':
            healthy_range = (0.0, 0.0)
            if '–' in lower or '-' in lower:
                nums = [float(x) for x in re.findall(r'[\d\.]+', lower)]
                if len(nums) == 2:
                    warning_range = (nums[0], nums[1])
            else:
                warning_range = (1.0, 3.0)

        elif upper == '100%':
            healthy_range = (99.8, 100.0)
            warning_range = (99.0, 99.8)

        kpi_rules[col_name] = {
            'kpi_name': name,
            'upper_text': upper,
            'lower_text': lower,
            'healthy_range': healthy_range,
            'warning_range': warning_range
        }

    # Custom overrides based on operational feedback
    kpi_rules['total_number_of_wifi_connected_subs'] = {
        'kpi_name': 'Total Number of WiFi Connected Subs',
        'upper_text': '~100',
        'lower_text': '<80',
        'healthy_range': (95.0, 105.0),
        'warning_range': (80.0, 95.0)
    }
    kpi_rules['total_number_of_lte_connected_subs'] = {
        'kpi_name': 'Total Number of LTE Connected Subs',
        'upper_text': '~2790',
        'lower_text': '<2500',
        'healthy_range': (2700.0, 2880.0),
        'warning_range': (2500.0, 2700.0)
    }
    kpi_rules['online_users'] = {
        'kpi_name': 'Online Users',
        'upper_text': '~1100',
        'lower_text': '<950',
        'healthy_range': (1050.0, 1150.0),
        'warning_range': (950.0, 1050.0)
    }

    return kpi_rules


if __name__ == "__main__":
    rules = parse_limit_rules()
    print(f"Loaded {len(rules)} KPI limit rules from Excel:")
    for c, r in list(rules.items())[:10]:
        print(f"  {c:30s} -> Healthy: {r['healthy_range']} | Warning: {r['warning_range']}")
